import io, os, sys, tempfile, unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pypdf import PdfReader
import pypdfium2 as pdfium
from wattanapat_pdf.engine import ops
from wattanapat_pdf.engine.ops import EngineError, PasswordRequired, WrongPassword
from helpers import make_pdf, make_image


class Base(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.d = Path(self._td.name)
        ops.set_work_dir(self.d / "work")
        self.out = self.d / "out"

    def tearDown(self):
        self._td.cleanup()

    def pdf(self, name="a.pdf", **kw):
        return make_pdf(self.d / name, **kw)


class TestRanges(unittest.TestCase):
    def test_groups(self):
        g = ops.parse_range_groups("1-5, 8, 12-15", 20)
        self.assertEqual([x[0] for x in g], ["1-5", "8", "12-15"])
        self.assertEqual(g[0][1], [0, 1, 2, 3, 4])
        self.assertEqual(ops.parse_ranges("18-", 20), [17, 18, 19])

    def test_errors(self):
        for bad in ("", "abc", "0-3", "5-2", "1-99"):
            with self.assertRaises(EngineError, msg=bad):
                ops.parse_ranges(bad, 10)


class TestMergeSplitOrganize(Base):
    def test_merge_with_ranges_and_order(self):
        a = self.pdf("a.pdf", pages=3, label="A")
        b = self.pdf("b.pdf", pages=4, label="B")
        out = ops.merge([{"path": b, "pages": "2-3"}, {"path": a}], self.out / "m.pdf")
        r = PdfReader(out)
        self.assertEqual(len(r.pages), 5)
        self.assertIn("B 2 of 4", r.pages[0].extract_text())
        self.assertIn("A 1 of 3", r.pages[2].extract_text())

    def test_never_overwrites(self):
        a = self.pdf("a.pdf")
        o1 = ops.merge([{"path": a}], self.out / "m.pdf")
        o2 = ops.merge([{"path": a}], self.out / "m.pdf")
        self.assertNotEqual(o1, o2)
        self.assertTrue(Path(o1).exists() and Path(o2).exists())

    def test_split_ranges_and_each(self):
        a = self.pdf("doc.pdf", pages=10)
        outs = ops.split(a, self.out, "ranges", "1-5, 8, 10")
        self.assertEqual([len(PdfReader(o).pages) for o in outs], [5, 1, 1])
        self.assertEqual([Path(o).name for o in outs], ["doc_p1-5.pdf", "doc_p8.pdf", "doc_p10.pdf"])
        outs = ops.split(a, self.d / "each", "each")
        self.assertEqual(len(outs), 10)
        self.assertEqual(Path(outs[0]).name, "doc_p01.pdf")

    def test_organize_reorder_rotate_delete(self):
        a = self.pdf("a.pdf", pages=4)
        out = ops.organize(a, [[3, 0], [0, 90], [1, 0]], self.out / "o.pdf")
        r = PdfReader(out)
        self.assertEqual(len(r.pages), 3)
        self.assertIn("Page 4", r.pages[0].extract_text())
        self.assertEqual(r.pages[1].rotation, 90)
        self.assertIn("Page 2", r.pages[2].extract_text())
        with self.assertRaises(EngineError):
            ops.organize(a, [], self.out / "x.pdf")


class TestImages(Base):
    def test_images_to_pdf_a4_and_fit(self):
        i1 = make_image(self.d / "1.jpg")
        i2 = make_image(self.d / "2.png", size=(300, 900), fmt="PNG")
        out = ops.images_to_pdf([i1, i2], self.out / "i.pdf", "A4")
        r = PdfReader(out)
        self.assertEqual(len(r.pages), 2)
        self.assertAlmostEqual(float(r.pages[0].mediabox.width), 841.89, delta=1)   # landscape A4
        self.assertAlmostEqual(float(r.pages[1].mediabox.width), 595.28, delta=1)   # portrait A4
        out = ops.images_to_pdf([i1], self.out / "f.pdf", "fit")
        self.assertAlmostEqual(float(PdfReader(out).pages[0].mediabox.width), 800 * 72 / 150, delta=1)

    def test_pdf_to_images(self):
        a = self.pdf("a.pdf", pages=3)
        outs = ops.pdf_to_images(a, self.out, "png", 72, "2-3")
        self.assertEqual([Path(o).name for o in outs], ["a_p2.png", "a_p3.png"])
        outs = ops.pdf_to_images(a, self.out, "jpeg", 72)
        self.assertTrue(all(o.endswith(".jpg") for o in outs))


class TestProtect(Base):
    def test_user_password_roundtrip_and_permissions(self):
        a = self.pdf("a.pdf")
        out = ops.protect(a, self.out / "p.pdf", "open123", "owner456", allow_print=True,
                          allow_modify=False, allow_copy=False)
        r = PdfReader(out)
        self.assertTrue(r.is_encrypted)
        self.assertEqual(r.decrypt("wrong"), 0)
        self.assertNotEqual(r.decrypt("open123"), 0)
        perms = r.user_access_permissions
        from pypdf.constants import UserAccessPermissions as U
        self.assertTrue(perms & U.PRINT)
        self.assertFalse(perms & U.MODIFY)
        self.assertFalse(perms & U.EXTRACT)
        self.assertIn("Page 1", r.pages[0].extract_text())
        raw = Path(out).read_bytes()
        self.assertNotIn(b"open123", raw)
        self.assertIn(b"/V 5", raw)  # AES-256

    def test_owner_only_opens_without_password(self):
        a = self.pdf("a.pdf")
        out = ops.protect(a, self.out / "p.pdf", "", "", allow_copy=False)
        r = PdfReader(out)
        self.assertTrue(r.is_encrypted)
        self.assertNotEqual(r.decrypt(""), 0)

    def test_requires_something(self):
        with self.assertRaises(EngineError):
            ops.protect(self.pdf("a.pdf"), self.out / "p.pdf")

    def test_password_required_and_wrong(self):
        a = self.pdf("a.pdf")
        enc = ops.protect(a, self.out / "p.pdf", "pw", "own")
        with self.assertRaises(PasswordRequired):
            ops.open_reader(enc)
        with self.assertRaises(WrongPassword):
            ops.open_reader(enc, "nope")
        self.assertTrue(ops.inspect(enc)["needs_password"])
        self.assertEqual(ops.inspect(enc, "pw")["pages"], 3)
        out = ops.merge([{"path": enc, "password": "pw"}], self.out / "m.pdf")
        self.assertEqual(len(PdfReader(out).pages), 3)


class TestWatermark(Base):
    SPEC = {"kind": "text", "text": "CONFIDENTIAL", "position": 4, "opacity": 0.5,
            "font_size": 60, "color": "#ff0000", "rotation": 45, "font": "auto"}

    def _red_pixels(self, path, page=0):
        d = pdfium.PdfDocument(path)
        im = d[page].render(scale=0.5).to_pil().convert("RGB")
        d.close()
        return sum(1 for r, g, b in im.getdata() if r > 200 and g < 150 and b < 150)

    def test_text_watermark_all_pages(self):
        a = self.pdf("a.pdf", pages=2)
        out = ops.watermark(a, self.out / "w.pdf", self.SPEC)
        self.assertGreater(self._red_pixels(out, 0), 50)
        self.assertGreater(self._red_pixels(out, 1), 50)
        self.assertEqual(self._red_pixels(a, 0), 0)

    def test_page_selection_and_all_nine_positions(self):
        a = self.pdf("a.pdf", pages=2)
        out = ops.watermark(a, self.out / "w.pdf", dict(self.SPEC, pages="2"))
        self.assertEqual(self._red_pixels(out, 0), 0)
        self.assertGreater(self._red_pixels(out, 1), 50)
        for pos in range(9):
            o = ops.watermark(a, self.out / f"p{pos}.pdf", dict(self.SPEC, position=pos, rotation=0, font_size=24))
            self.assertGreater(self._red_pixels(o), 10, pos)

    def test_rotated_page_and_letter_size(self):
        a = self.pdf("a.pdf", pages=1, size=(612, 792))
        from pypdf import PdfWriter
        r = PdfReader(a); w = PdfWriter(); w.add_page(r.pages[0]); w.pages[0].rotate(90)
        rp = self.d / "rot.pdf"; w.write(str(rp))
        out = ops.watermark(str(rp), self.out / "w.pdf", self.SPEC)
        self.assertGreater(self._red_pixels(out), 50)

    def test_image_watermark(self):
        a = self.pdf("a.pdf", pages=1)
        png = make_image(self.d / "stamp.png", (200, 100), (255, 0, 0), "PNG")
        out = ops.watermark(a, self.out / "w.pdf", {"kind": "image", "image_path": png, "position": 8,
                                                    "opacity": 0.6, "rotation": 0, "image_scale": 0.3})
        self.assertGreater(self._red_pixels(out), 100)

    def test_preview_and_errors(self):
        a = self.pdf("a.pdf", pages=2)
        pv = ops.watermark_preview(a, self.SPEC, 300)
        self.assertEqual(pv["data"][:2], b"\xff\xd8")
        with self.assertRaises(EngineError):
            ops.watermark(a, self.out / "x.pdf", dict(self.SPEC, text="  "))
        with self.assertRaises(EngineError):  # no font on this box can draw Thai -> clear error, not garbage
            from wattanapat_pdf.engine import fonts
            if fonts.pick("ก", "auto").covers:
                raise EngineError("thai font present")
            ops.watermark(a, self.out / "x.pdf", dict(self.SPEC, text="เอกสารภายใน"))


class TestCompress(Base):
    def test_levels_reduce_size_and_stay_valid(self):
        a = self.pdf("img.pdf", pages=3, with_image=True)
        orig = Path(a).stat().st_size
        sizes = {}
        for lvl in ("quality", "email", "smallest"):
            res = ops.compress(a, self.out / f"{lvl}.pdf", lvl)
            sizes[lvl] = res["new"]
            self.assertEqual(len(PdfReader(res["path"]).pages), 3)
            self.assertLessEqual(res["new"], orig)
            est = ops.estimate_compress(a, lvl)
            self.assertAlmostEqual(est["estimated"], res["new"], delta=max(2000, res["new"] * 0.05))
        self.assertLess(sizes["smallest"], sizes["email"])
        self.assertLess(sizes["email"], orig * 0.6)

    def test_never_bigger(self):
        a = self.pdf("t.pdf", pages=2)
        res = ops.compress(a, self.out / "c.pdf", "quality")
        self.assertLessEqual(res["new"], Path(a).stat().st_size)


class TestConvert(Base):
    def test_pdf_to_docx_and_back(self):
        a = self.pdf("a.pdf", pages=2)
        res = ops.pdf_to_docx(a, self.out / "a.docx")
        from docx import Document
        txt = "\n".join(p.text for p in Document(res["path"]).paragraphs)
        self.assertIn("Page 1 of 2", txt)
        self.assertIn("Page 2 of 2", txt)
        self.assertTrue(res["warnings"])
        back = ops.docx_to_pdf(res["path"], self.out / "back.pdf")
        t = "".join(p.extract_text() for p in PdfReader(back["path"]).pages)
        self.assertIn("Page 1 of 2", t)

    def test_docx_to_pdf_table_heading_image(self):
        from docx import Document
        d = Document()
        d.add_heading("Report <&> Title", 1)
        p = d.add_paragraph(); p.add_run("bold ").bold = True; p.add_run("plain")
        t = d.add_table(rows=2, cols=3)
        for r in range(2):
            for c in range(3):
                t.cell(r, c).text = f"r{r}c{c}"
        d.add_picture(make_image(self.d / "i.png", (300, 200), fmt="PNG"))
        src = self.d / "t.docx"; d.save(str(src))
        res = ops.docx_to_pdf(str(src), self.out / "t.pdf")
        text = PdfReader(res["path"]).pages[0].extract_text()
        for s in ("Report", "bold plain", "r1c2"):
            self.assertIn(s, text)

    def test_scanned_page_becomes_picture(self):
        img = make_image(self.d / "s.jpg")
        scan = ops.images_to_pdf([img], self.out / "scan.pdf", "A4")
        res = ops.pdf_to_docx(scan, self.out / "scan.docx")
        self.assertTrue(any("ภาพสแกน" in w for w in res["warnings"]))


class TestRenderAndErrors(Base):
    def test_render_pages(self):
        a = self.pdf("a.pdf", pages=3)
        r = ops.render_pages(a, [0, 2], 160)
        self.assertEqual([x["index"] for x in r], [0, 2])
        self.assertEqual(r[0]["w"], 160)

    def test_corrupt_and_missing(self):
        bad = self.d / "bad.pdf"; bad.write_bytes(b"not a pdf at all")
        with self.assertRaises(EngineError):
            ops.open_reader(bad)
        with self.assertRaises(EngineError):
            ops.open_reader(self.d / "missing.pdf")

    def test_originals_untouched(self):
        a = self.pdf("a.pdf", pages=3)
        before = Path(a).read_bytes()
        ops.watermark(a, self.out / "w.pdf", TestWatermark.SPEC)
        ops.compress(a, self.out / "c.pdf", "email")
        ops.organize(a, [[0, 90]], self.out / "o.pdf")
        ops.protect(a, self.out / "p.pdf", "x", "y")
        self.assertEqual(before, Path(a).read_bytes())


class TestIPC(Base):
    def test_engine_process_roundtrip_progress_errors_cancel(self):
        from wattanapat_pdf.engine.service import EngineClient, Cancelled
        c = EngineClient(str(self.d / "work"))
        try:
            a = self.pdf("a.pdf", pages=5)
            seen = []
            outs = c.call("split", progress=lambda f, t: seen.append(f), src=a, out_dir=str(self.out), mode="each")
            self.assertEqual(len(outs), 5)
            self.assertTrue(seen and seen[-1] == 1.0)
            with self.assertRaises(EngineError):
                c.call("split", src=a, out_dir=str(self.out), mode="ranges", ranges_text="99")
            enc = c.call("protect", src=a, out_path=str(self.out / "e.pdf"), user_password="z", owner_password="y")
            with self.assertRaises(PasswordRequired):
                c.call("merge", items=[{"path": enc}], out_path=str(self.out / "m.pdf"))
            self.assertEqual(c.call("inspect", path=a)["pages"], 5)  # engine still healthy after errors
            # cancel + automatic restart
            big = self.pdf("big.pdf", pages=60, with_image=True)
            import threading, time
            threading.Timer(0.4, c.cancel).start()
            with self.assertRaises(Cancelled):
                c.call("compress", src=big, out_path=str(self.out / "big_c.pdf"), level="smallest")
            self.assertEqual(c.call("inspect", path=a)["pages"], 5)
        finally:
            c.stop()


class TestSettingsSandboxLog(Base):
    def test_presets_and_recent(self):
        from wattanapat_pdf.settings import Settings
        s = Settings(self.d / "s.json")
        self.assertEqual([p["name"] for p in s.presets], ["ชุดที่ 1", "ชุดที่ 2", "ชุดที่ 3", "ชุดที่ 4"])
        i = s.add_preset("ส่งลูกค้า", s.presets[0]); s.rename_preset(0, "เอกสารภายใน")
        self.assertEqual(Settings(self.d / "s.json").presets[0]["name"], "เอกสารภายใน")
        s.delete_preset(i); self.assertEqual(len(s.presets), 4)
        s.add_recent("C:/x.pdf"); self.assertEqual(s.get("recent"), ["C:/x.pdf"])
        s.set("remember_recent", False); s.add_recent("C:/y.pdf"); self.assertEqual(len(s.get("recent")), 1)
        s.clear_history(); self.assertEqual(s.get("recent"), [])
        self.assertNotIn("password", (self.d / "s.json").read_text().lower())

    def test_sandbox_lifecycle(self):
        from wattanapat_pdf.sandbox import Sandbox
        old = Sandbox(self.d / "sb"); old.start()
        stale = old.root / "session-999999-dead"; stale.mkdir(); (stale / "x.tmp").write_text("x")
        sb = Sandbox(self.d / "sb"); sb.start()
        self.assertFalse(stale.exists())
        (sb.session / "a.tmp").write_bytes(b"123"); (sb.session / "d").mkdir(); (sb.session / "d" / "b").write_bytes(b"4567")
        self.assertEqual(sb.stats(), (2, 7))
        self.assertEqual(sb.clear(), 2); self.assertEqual(sb.stats(), (0, 0))
        sb.close(); self.assertFalse(sb.session.exists()); old.close()

    def test_log_has_no_filenames_or_passwords(self):
        os.environ["WATTANAPAT_DATA_DIR"] = str(self.d / "data")
        from wattanapat_pdf import applog
        applog._ready = False; applog._log.handlers.clear()
        try:
            open("/nonexistent/SECRET-client-contract.pdf")
        except OSError as e:
            applog.log_exception("merge", e)
        txt = applog.log_file().read_text(encoding="utf-8")
        self.assertIn("op=merge exc=FileNotFoundError", txt)
        self.assertNotIn("SECRET", txt)

    def test_updater_https_only(self):
        from wattanapat_pdf import updater
        with self.assertRaises(ValueError):
            updater.check("http://example.com/v.json")
        self.assertTrue(updater._ver("1.2.0") > updater._ver("1.1.9"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
