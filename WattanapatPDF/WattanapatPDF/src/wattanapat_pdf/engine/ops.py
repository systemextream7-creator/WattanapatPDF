"""PDF processing operations.  100% local - no network access anywhere in this module.

Every function writes to a private work folder first and then moves the result to a
*new* file name; an existing file is never overwritten.
"""
from __future__ import annotations

import io, os, re, secrets, shutil, tempfile, uuid
from pathlib import Path
from typing import Callable, Optional

from pypdf import PdfReader, PdfWriter, Transformation
from pypdf.errors import PdfReadError, PdfStreamError, EmptyFileError

Progress = Optional[Callable[[float, str], None]]
_WORK = Path(tempfile.gettempdir()) / "wattanapat-engine"


class EngineError(Exception):
    """Error whose message is safe and meaningful to show to the user (Thai)."""


class PasswordRequired(EngineError):
    pass


class WrongPassword(EngineError):
    pass


def set_work_dir(p) -> None:
    global _WORK
    _WORK = Path(p)
    _WORK.mkdir(parents=True, exist_ok=True)


def _tmp(suffix: str) -> Path:
    _WORK.mkdir(parents=True, exist_ok=True)
    return _WORK / f"{uuid.uuid4().hex}{suffix}"


def _p(progress: Progress, f: float, text: str = "") -> None:
    if progress:
        progress(max(0.0, min(1.0, f)), text)


def unique_path(path) -> Path:
    path = Path(path)
    if not path.exists():
        return path
    for n in range(1, 10000):
        cand = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not cand.exists():
            return cand
    raise EngineError("ไม่สามารถตั้งชื่อไฟล์ปลายทางได้")


def _finalize(tmp: Path, dest) -> str:
    dest = Path(dest)
    try:
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest = unique_path(dest)
        shutil.move(str(tmp), str(dest))
    except OSError as e:
        raise EngineError("บันทึกไฟล์ไม่สำเร็จ ตรวจสอบสิทธิ์และพื้นที่ว่างของโฟลเดอร์ปลายทาง") from e
    return str(dest)


# ---------------------------------------------------------------- reading --
def open_reader(path, password: str | None = None) -> PdfReader:
    try:
        r = PdfReader(str(path))
    except FileNotFoundError as e:
        raise EngineError("ไม่พบไฟล์ต้นฉบับ") from e
    except (PdfReadError, PdfStreamError, EmptyFileError, ValueError, KeyError, TypeError) as e:
        raise EngineError("ไฟล์ PDF เสียหายหรืออ่านไม่ได้") from e
    except OSError as e:
        raise EngineError("เปิดไฟล์ไม่ได้ (ไฟล์อาจถูกใช้งานอยู่หรือไม่มีสิทธิ์เข้าถึง)") from e
    if r.is_encrypted:
        if r.decrypt(password or "") == 0:
            raise (WrongPassword("รหัสผ่านไม่ถูกต้อง") if password else PasswordRequired("ไฟล์นี้มีรหัสผ่าน"))
    return r


def inspect(path, password: str | None = None, progress: Progress = None) -> dict:
    size = Path(path).stat().st_size if Path(path).exists() else 0
    try:
        r = open_reader(path, password)
    except PasswordRequired:
        return {"pages": 0, "size": size, "needs_password": True}
    return {"pages": len(r.pages), "size": size, "needs_password": False}


def parse_range_groups(text: str, total: int) -> list[tuple[str, list[int]]]:
    """'1-5, 8, 12-' -> [('1-5',[0..4]), ('8',[7]), ('12-N',[...])]  (zero-based indexes)."""
    text = (text or "").strip()
    if not text:
        raise EngineError("กรุณาระบุช่วงหน้า เช่น 1-5, 8, 12-15")
    groups = []
    for tok in [t for t in re.split(r"[,;\s]+", text) if t]:
        m = re.fullmatch(r"(\d+)(?:-(\d*))?", tok)
        if not m:
            raise EngineError(f"ช่วงหน้า '{tok}' ไม่ถูกต้อง (ตัวอย่างที่ถูก: 1-5, 8, 12-15)")
        a = int(m.group(1))
        if m.group(2) is None:
            b = a
        else:
            b = int(m.group(2)) if m.group(2) else total
        if a < 1 or b < a:
            raise EngineError(f"ช่วงหน้า '{tok}' ไม่ถูกต้อง")
        if b > total:
            raise EngineError(f"หน้า {b} เกินจำนวนหน้าทั้งหมด ({total} หน้า)")
        groups.append((f"{a}" if a == b else f"{a}-{b}", list(range(a - 1, b))))
    return groups


def parse_ranges(text: str, total: int) -> list[int]:
    return [i for _, g in parse_range_groups(text, total) for i in g]


# ------------------------------------------------------------------ merge --
def merge(items: list[dict], out_path: str, progress: Progress = None) -> str:
    """items: [{path, password?, pages?}] -> one PDF."""
    if not items:
        raise EngineError("ยังไม่ได้เลือกไฟล์ที่จะรวม")
    w = PdfWriter()
    readers = []
    for n, it in enumerate(items):
        _p(progress, n / (len(items) + 1), "กำลังอ่านไฟล์")
        r = open_reader(it["path"], it.get("password"))
        readers.append(r)
        idx = parse_ranges(it["pages"], len(r.pages)) if (it.get("pages") or "").strip() else range(len(r.pages))
        for i in idx:
            w.add_page(r.pages[i])
    _p(progress, 0.9, "กำลังบันทึก")
    tmp = _tmp(".pdf")
    with open(tmp, "wb") as f:
        w.write(f)
    _p(progress, 1.0, "เสร็จสิ้น")
    return _finalize(tmp, out_path)


# ------------------------------------------------------------------ split --
def split(src: str, out_dir: str, mode: str = "ranges", ranges_text: str = "",
          password: str | None = None, progress: Progress = None) -> list[str]:
    """mode 'ranges': one file per comma separated range; mode 'each': one file per page."""
    r = open_reader(src, password)
    total = len(r.pages)
    stem = Path(src).stem
    if mode == "each":
        width = len(str(total))
        groups = [(f"{i + 1:0{width}d}", [i]) for i in range(total)]
    else:
        groups = parse_range_groups(ranges_text, total)
    outs = []
    for n, (label, idx) in enumerate(groups):
        _p(progress, n / len(groups), f"กำลังแยก {n + 1}/{len(groups)}")
        w = PdfWriter()
        for i in idx:
            w.add_page(r.pages[i])
        tmp = _tmp(".pdf")
        with open(tmp, "wb") as f:
            w.write(f)
        outs.append(_finalize(tmp, Path(out_dir) / f"{stem}_p{label}.pdf"))
    _p(progress, 1.0, "เสร็จสิ้น")
    return outs


# --------------------------------------------------------------- organize --
def organize(src: str, pages: list[list[int]], out_path: str, password: str | None = None,
             progress: Progress = None) -> str:
    """pages: ordered [[source_index, extra_rotation_deg], ...]; omitted pages are deleted."""
    if not pages:
        raise EngineError("ไม่มีหน้าเหลือให้บันทึก")
    r = open_reader(src, password)
    w = PdfWriter()
    for n, (i, rot) in enumerate(pages):
        _p(progress, n / len(pages), "กำลังจัดเรียงหน้า")
        w.add_page(r.pages[i])
        if rot % 360:
            w.pages[-1].rotate(rot % 360)
    tmp = _tmp(".pdf")
    with open(tmp, "wb") as f:
        w.write(f)
    _p(progress, 1.0, "เสร็จสิ้น")
    return _finalize(tmp, out_path)


# ----------------------------------------------------------- images -> PDF --
def images_to_pdf(images: list[str], out_path: str, page_size: str = "A4", margin_mm: float = 0,
                  progress: Progress = None) -> str:
    from PIL import Image, ImageOps, ImageSequence
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    if not images:
        raise EngineError("ยังไม่ได้เลือกรูปภาพ")
    tmp = _tmp(".pdf")
    c = canvas.Canvas(str(tmp))
    for n, path in enumerate(images):
        _p(progress, n / len(images), "กำลังนำเข้ารูปภาพ")
        try:
            im0 = Image.open(path)
        except Exception as e:
            raise EngineError("เปิดไฟล์รูปภาพไม่ได้ หรือรูปแบบไม่รองรับ") from e
        is_jpeg = (im0.format == "JPEG")
        dpi = im0.info.get("dpi", (150, 150))[0] or 150
        dpi = min(max(float(dpi), 50.0), 600.0)
        for frame in ImageSequence.Iterator(im0):
            rotated = frame.getexif().get(0x0112, 1) not in (0, 1)  # EXIF orientation needs re-encoding
            im = ImageOps.exif_transpose(frame.copy())
            if im.mode in ("RGBA", "LA", "P", "PA"):
                im = im.convert("RGBA")
                bg = Image.new("RGB", im.size, "white")
                bg.paste(im, mask=im.split()[-1])
                im, is_jpeg = bg, False
            elif im.mode not in ("RGB", "L"):
                im, is_jpeg = im.convert("RGB"), False
            iw, ih = im.size
            if page_size == "fit":
                pw, ph = iw * 72.0 / dpi, ih * 72.0 / dpi
                c.setPageSize((pw, ph))
                x, y, dw, dh = 0, 0, pw, ph
            else:
                pw, ph = (A4[1], A4[0]) if iw > ih else A4
                c.setPageSize((pw, ph))
                m = margin_mm * mm
                s = min((pw - 2 * m) / iw, (ph - 2 * m) / ih)
                dw, dh = iw * s, ih * s
                x, y = (pw - dw) / 2, (ph - dh) / 2
            src = path if (is_jpeg and not rotated) else im  # untouched JPEGs are embedded as-is
            c.drawImage(ImageReader(src), x, y, dw, dh)
            c.showPage()
        im0.close()
    c.save()
    _p(progress, 1.0, "เสร็จสิ้น")
    return _finalize(tmp, out_path)


# ----------------------------------------------------------- PDF -> images --
def pdf_to_images(src: str, out_dir: str, fmt: str = "png", dpi: int = 150, pages_text: str = "",
                  password: str | None = None, progress: Progress = None) -> list[str]:
    import pypdfium2 as pdfium
    try:
        doc = pdfium.PdfDocument(str(src), password=password or None)
    except pdfium.PdfiumError as e:
        msg = str(e).lower()
        if "password" in msg:
            raise (WrongPassword("รหัสผ่านไม่ถูกต้อง") if password else PasswordRequired("ไฟล์นี้มีรหัสผ่าน")) from e
        raise EngineError("ไฟล์ PDF เสียหายหรืออ่านไม่ได้") from e
    try:
        total = len(doc)
        idx = parse_ranges(pages_text, total) if (pages_text or "").strip() else list(range(total))
        stem, width, outs = Path(src).stem, len(str(total)), []
        ext = "jpg" if fmt.lower() in ("jpg", "jpeg") else "png"
        for n, i in enumerate(idx):
            _p(progress, n / len(idx), f"กำลังแปลงหน้า {i + 1}")
            img = doc[i].render(scale=dpi / 72.0).to_pil()
            tmp = _tmp("." + ext)
            if ext == "jpg":
                img.convert("RGB").save(tmp, "JPEG", quality=90, optimize=True)
            else:
                img.save(tmp, "PNG", optimize=True)
            outs.append(_finalize(tmp, Path(out_dir) / f"{stem}_p{i + 1:0{width}d}.{ext}"))
    finally:
        doc.close()
    _p(progress, 1.0, "เสร็จสิ้น")
    return outs


# -------------------------------------------------------------- protection --
def protect(src: str, out_path: str, user_password: str = "", owner_password: str = "",
            allow_print: bool = True, allow_modify: bool = True, allow_copy: bool = True,
            password: str | None = None, progress: Progress = None) -> str:
    restricted = not (allow_print and allow_modify and allow_copy)
    if not user_password and not owner_password and not restricted:
        raise EngineError("กรุณาตั้งรหัสผ่านอย่างน้อยหนึ่งรหัส หรือเลือกจำกัดสิทธิ์")
    if not owner_password:
        owner_password = secrets.token_urlsafe(24)  # restrictions need an owner secret; never stored
    r = open_reader(src, password)
    _p(progress, 0.3, "กำลังเข้ารหัส")
    w = PdfWriter(clone_from=r)
    flags = -1  # all permissions granted
    if not allow_print:
        flags &= ~(4 | 2048)
    if not allow_modify:
        flags &= ~(8 | 32 | 256 | 1024)
    if not allow_copy:
        flags &= ~16
    w.encrypt(user_password=user_password, owner_password=owner_password,
              permissions_flag=flags, algorithm="AES-256")
    tmp = _tmp(".pdf")
    with open(tmp, "wb") as f:
        w.write(f)
    _p(progress, 1.0, "เสร็จสิ้น")
    return _finalize(tmp, out_path)


# --------------------------------------------------------------- watermark --
def _hex_to_rgb(h: str):
    h = (h or "#808080").lstrip("#")
    if len(h) != 6:
        h = "808080"
    return tuple(int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))


def _overlay_pdf(w: float, h: float, spec: dict) -> PdfReader:
    import math
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.lib.utils import ImageReader
    from . import fonts

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(w, h))
    pos = int(spec.get("position", 4))
    col, row = pos % 3, pos // 3
    rot = float(spec.get("rotation", 0))
    c.saveState()
    c.setFillAlpha(max(0.02, min(1.0, float(spec.get("opacity", 0.3)))))
    if spec.get("kind") == "image":
        try:
            ir = ImageReader(spec["image_path"])
        except Exception as e:
            raise EngineError("เปิดรูปภาพลายน้ำไม่ได้") from e
        iw, ih = ir.getSize()
        bw = w * float(spec.get("image_scale", 0.4))
        bh = bw * ih / iw
        draw = lambda: c.drawImage(ir, -bw / 2, -bh / 2, bw, bh, mask="auto")
    else:
        text = spec.get("text", "")
        if not text.strip():
            raise EngineError("กรุณาพิมพ์ข้อความลายน้ำ")
        size = float(spec.get("font_size", 48))
        fc = fonts.pick(text, spec.get("font", "auto"))
        if not fc.covers:
            raise EngineError("ไม่พบฟอนต์ที่รองรับตัวอักษรในข้อความลายน้ำ กรุณาเลือกฟอนต์อื่น")
        bw, bh = stringWidth(text, fc.name, size), size
        r_, g_, b_ = _hex_to_rgb(spec.get("color", "#808080"))
        c.setFillColorRGB(r_, g_, b_)

        def draw():
            c.setFont(fc.name, size)
            c.drawCentredString(0, -size * 0.3, text)
    a = math.radians(rot)
    hw = (abs(math.cos(a)) * bw + abs(math.sin(a)) * bh) / 2
    hh = (abs(math.sin(a)) * bw + abs(math.cos(a)) * bh) / 2
    m = 24.0
    ax = {0: m + hw, 1: w / 2, 2: w - m - hw}[col]
    ay = {0: h - m - hh, 1: h / 2, 2: m + hh}[row]
    c.translate(ax, ay)
    c.rotate(rot)
    draw()
    c.restoreState()
    c.showPage()
    c.save()
    buf.seek(0)
    return PdfReader(buf)


def _apply_watermark(reader: PdfReader, spec: dict, only_first: bool = False,
                     progress: Progress = None) -> PdfWriter:
    total = len(reader.pages)
    idx = {0} if only_first else (set(parse_ranges(spec["pages"], total)) if (spec.get("pages") or "").strip()
                                  else set(range(total)))
    w = PdfWriter()
    cache: dict[tuple, PdfReader] = {}
    for i in range(1 if only_first else total):
        _p(progress, i / total, "กำลังประทับลายน้ำ")
        page = w.add_page(reader.pages[i])
        if i not in idx:
            continue
        if int(page.get("/Rotate", 0) or 0) % 360:
            page.transfer_rotation_to_content()
        mb = page.mediabox
        pw, ph = float(mb.width), float(mb.height)
        key = (round(pw, 1), round(ph, 1))
        if key not in cache:
            cache[key] = _overlay_pdf(pw, ph, spec)
        page.merge_transformed_page(cache[key].pages[0],
                                    Transformation().translate(float(mb.left), float(mb.bottom)))
    return w


def watermark(src: str, out_path: str, spec: dict, password: str | None = None,
              progress: Progress = None) -> str:
    r = open_reader(src, password)
    w = _apply_watermark(r, spec, progress=progress)
    tmp = _tmp(".pdf")
    with open(tmp, "wb") as f:
        w.write(f)
    _p(progress, 1.0, "เสร็จสิ้น")
    return _finalize(tmp, out_path)


# -------------------------------------------------------------- compression --
LEVELS = {  # quality, longest image edge in px
    "quality": (85, None),
    "email": (65, 1800),
    "smallest": (40, 1100),
}


def _compress_into(reader: PdfReader, level: str, sink, progress: Progress = None) -> None:
    from PIL import Image
    q, max_edge = LEVELS[level]
    w = PdfWriter(clone_from=reader)
    pages = list(w.pages)
    for n, page in enumerate(pages):
        _p(progress, n / max(1, len(pages)) * 0.9, f"กำลังบีบอัดหน้า {n + 1}/{len(pages)}")
        try:
            images = list(page.images)
        except Exception:
            images = []
        for img in images:
            try:
                obj = img.indirect_reference.get_object()
                if "/SMask" in obj or "/Mask" in obj or obj.get("/ImageMask"):
                    continue
                raw_len = len(obj._data) if hasattr(obj, "_data") else 0
                if raw_len and raw_len < 20_000:
                    continue
                pil = img.image
                if pil.mode not in ("RGB", "L"):
                    if pil.mode in ("CMYK", "P", "1"):
                        pil = pil.convert("RGB")
                    else:
                        continue
                if max_edge and max(pil.size) > max_edge:
                    s = max_edge / max(pil.size)
                    pil = pil.resize((max(1, int(pil.width * s)), max(1, int(pil.height * s))), Image.LANCZOS)
                probe = io.BytesIO()
                pil.save(probe, "JPEG", quality=q, optimize=True)
                if raw_len and probe.tell() >= raw_len * 0.9:
                    continue
                img.replace(pil, quality=q)
            except Exception:
                continue  # never fail the whole file because one image is unusual
    w.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    for page in w.pages:
        try:
            page.compress_content_streams()
        except Exception:
            pass
    w.write(sink)


def estimate_compress(src: str, level: str, password: str | None = None, progress: Progress = None) -> dict:
    """Run the real compression in memory and report the resulting size (nothing is written to disk)."""
    r = open_reader(src, password)
    buf = io.BytesIO()
    _compress_into(r, level, buf, progress)
    orig = Path(src).stat().st_size
    new = min(buf.tell(), orig)
    return {"original": orig, "estimated": new}


def compress(src: str, out_path: str, level: str = "email", password: str | None = None,
             progress: Progress = None) -> dict:
    if level not in LEVELS:
        raise EngineError("ระดับการบีบอัดไม่ถูกต้อง")
    r = open_reader(src, password)
    tmp = _tmp(".pdf")
    with open(tmp, "wb") as f:
        _compress_into(r, level, f, progress)
    orig = Path(src).stat().st_size
    if tmp.stat().st_size >= orig and not r.is_encrypted:
        shutil.copyfile(src, tmp)  # never hand back a bigger file
    out = _finalize(tmp, out_path)
    _p(progress, 1.0, "เสร็จสิ้น")
    return {"path": out, "original": orig, "new": Path(out).stat().st_size}


# ------------------------------------------------------------------ preview --
def render_pages(path: str, indices: list[int], width: int = 200, password: str | None = None,
                 quality: int = 80, progress: Progress = None) -> list[dict]:
    """Render pages to JPEG bytes for thumbnails / previews: [{index, data, w, h, pw, ph}]."""
    import pypdfium2 as pdfium
    try:
        doc = pdfium.PdfDocument(str(path), password=password or None)
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise (WrongPassword("รหัสผ่านไม่ถูกต้อง") if password else PasswordRequired("ไฟล์นี้มีรหัสผ่าน")) from e
        raise EngineError("ไฟล์ PDF เสียหายหรืออ่านไม่ได้") from e
    out = []
    try:
        for n, i in enumerate(indices):
            _p(progress, n / max(1, len(indices)))
            pg = doc[i]
            pw, ph = pg.get_size()
            im = pg.render(scale=width / pw).to_pil().convert("RGB")
            b = io.BytesIO()
            im.save(b, "JPEG", quality=quality)
            out.append({"index": i, "data": b.getvalue(), "w": im.width, "h": im.height, "pw": pw, "ph": ph})
    finally:
        doc.close()
    return out


def watermark_preview(src: str, spec: dict, width: int = 600, password: str | None = None) -> dict:
    """Page 1 with the watermark applied, as JPEG bytes (never touches the disk)."""
    import pypdfium2 as pdfium
    r = open_reader(src, password)
    w = _apply_watermark(r, dict(spec, pages=""), only_first=True)
    buf = io.BytesIO()
    w.write(buf)
    doc = pdfium.PdfDocument(buf.getvalue())
    try:
        pg = doc[0]
        im = pg.render(scale=width / pg.get_size()[0]).to_pil().convert("RGB")
    finally:
        doc.close()
    b = io.BytesIO()
    im.save(b, "JPEG", quality=85)
    return {"data": b.getvalue(), "w": im.width, "h": im.height}


# ----------------------------------------------------------------- convert --
def pdf_to_docx(src: str, out_path: str, password: str | None = None, progress: Progress = None) -> dict:
    import pypdfium2 as pdfium
    from docx import Document
    from docx.shared import Cm
    from docx.oxml.ns import qn

    try:
        pdf = pdfium.PdfDocument(str(src), password=password or None)
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise (WrongPassword("รหัสผ่านไม่ถูกต้อง") if password else PasswordRequired("ไฟล์นี้มีรหัสผ่าน")) from e
        raise EngineError("ไฟล์ PDF เสียหายหรืออ่านไม่ได้") from e
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Tahoma"
    rpr = st.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:cs"), "Tahoma")
    scanned = 0
    try:
        n_pages = len(pdf)
        for i in range(n_pages):
            _p(progress, i / n_pages, f"กำลังแปลงหน้า {i + 1}/{n_pages}")
            page = pdf[i]
            text = page.get_textpage().get_text_range() or ""
            if len(text.strip()) < 15:  # scanned / image-only page: keep it visible as a picture
                scanned += 1
                buf = io.BytesIO()
                page.render(scale=150 / 72).to_pil().convert("RGB").save(buf, "PNG")
                buf.seek(0)
                doc.add_picture(buf, width=Cm(16))
            else:
                for line in text.splitlines():
                    if line.strip():
                        doc.add_paragraph(line.rstrip())
            if i < n_pages - 1:
                doc.add_page_break()
    finally:
        pdf.close()
    tmp = _tmp(".docx")
    doc.save(str(tmp))
    warnings = ["แปลงเฉพาะข้อความและลำดับบรรทัด การจัดหน้า ตาราง และคอลัมน์อาจไม่เหมือนต้นฉบับ"]
    if scanned:
        warnings.append(f"{scanned} หน้าเป็นภาพสแกน จึงใส่เป็นรูปภาพ (ไม่มี OCR จึงแก้ไขข้อความไม่ได้)")
    _p(progress, 1.0, "เสร็จสิ้น")
    return {"path": _finalize(tmp, out_path), "warnings": warnings}


def docx_to_pdf(src: str, out_path: str, progress: Progress = None) -> dict:
    from xml.sax.saxutils import escape
    from docx import Document
    from docx.table import Table as DTable
    from docx.text.paragraph import Paragraph as DPara
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from . import fonts

    try:
        d = Document(str(src))
    except Exception as e:
        raise EngineError("เปิดไฟล์ Word ไม่ได้ (รองรับเฉพาะ .docx)") from e
    reg, bold = fonts.family()
    body = ParagraphStyle("b", fontName=reg, fontSize=11, leading=17)
    cell = ParagraphStyle("c", parent=body, fontSize=9, leading=13)
    frame_w = A4[0] - 40 * mm
    warnings, skipped_img = [], 0

    def runs_markup(p) -> str:
        out = []
        for r in p.runs:
            t = escape(r.text or "").replace("\t", "    ")
            if not t:
                continue
            if r.bold:
                t = f"<b>{t}</b>"
            if r.italic:
                t = f"<i>{t}</i>"
            if r.underline:
                t = f"<u>{t}</u>"
            out.append(t)
        return "".join(out)

    story = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = DPara(child, d)
            style_name = (p.style.name or "") if p.style is not None else ""
            m = runs_markup(p)
            if style_name.startswith("List"):
                m = "• " + m
            if m.strip():
                if style_name == "Title" or style_name.startswith("Heading"):
                    lvl = 0 if style_name == "Title" else int(re.sub(r"\D", "", style_name) or 1)
                    size = max(12, 22 - 3 * lvl)
                    story.append(Paragraph(f"<b>{m}</b>", ParagraphStyle("h", parent=body, fontName=bold,
                                                                         fontSize=size, leading=size + 6,
                                                                         spaceBefore=8, spaceAfter=4)))
                else:
                    story.append(Paragraph(m, body))
            else:
                story.append(Spacer(1, 8))
            try:
                for rid in child.xpath(".//a:blip/@r:embed"):
                    blob = d.part.related_parts[rid].blob
                    from PIL import Image
                    iw, ih = Image.open(io.BytesIO(blob)).size
                    s = min(1.0, frame_w / iw, 600 / ih)
                    story.append(RLImage(io.BytesIO(blob), width=iw * s, height=ih * s))
            except Exception:
                skipped_img += 1
        elif tag == "tbl":
            t = DTable(child, d)
            data = [[Paragraph(escape(c.text).replace("\n", "<br/>"), cell) for c in row.cells] for row in t.rows]
            if data and data[0]:
                ncol = len(data[0])
                tbl = Table(data, colWidths=[frame_w / ncol] * ncol, repeatRows=1)
                tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                         ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                         ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F1FB"))]))
                story.append(tbl)
                story.append(Spacer(1, 8))
    if not story:
        raise EngineError("ไม่พบเนื้อหาในไฟล์ Word")
    if skipped_img:
        warnings.append(f"รูปภาพบางรายการนำเข้าไม่ได้ ({skipped_img})")
    warnings.append("แปลงแบบพื้นฐาน: ฟอนต์ สี ส่วนหัว/ท้ายกระดาษ และตารางซับซ้อนอาจต่างจากต้นฉบับ")
    tmp = _tmp(".pdf")
    SimpleDocTemplate(str(tmp), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=20 * mm, bottomMargin=20 * mm).build(story)
    _p(progress, 1.0, "เสร็จสิ้น")
    return {"path": _finalize(tmp, out_path), "warnings": warnings}


OPS = {
    "inspect": inspect, "merge": merge, "split": split, "organize": organize,
    "images_to_pdf": images_to_pdf, "pdf_to_images": pdf_to_images, "protect": protect,
    "watermark": watermark, "compress": compress, "estimate_compress": estimate_compress,
    "render_pages": render_pages, "watermark_preview": watermark_preview,
    "pdf_to_docx": pdf_to_docx, "docx_to_pdf": docx_to_pdf,
}
